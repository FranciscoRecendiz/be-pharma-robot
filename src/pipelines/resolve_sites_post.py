#!/usr/bin/env python3
"""
BePharma - Enriquecimiento de Websites 
Con validación multi-capa y filtros de calidad
"""

import asyncio
import pandas as pd
import aiohttp
from pathlib import Path
from datetime import datetime
from typing import List, Dict
from ddgs import DDGS

# Importa el resolver mejorado
from src.utils.site_resolver import (
    resolve_official_site,
    extract_domain,
    is_blocked_domain,
    normalize_company_name
)

# =============================================================================
# CONFIGURACIÓN
# =============================================================================

INPUT_FILE = Path("/project/data/processed/ecuador/empresas_raw.csv")
OUTPUT_FILE = Path("/project/data/processed/ecuador/empresas_final.csv")
REJECTED_FILE = Path("/project/data/processed/ecuador/empresas_rejected.csv")

Path(OUTPUT_FILE.parent).mkdir(parents=True, exist_ok=True)

# Configuración de procesamiento
BATCH_SIZE = 15  # Reducido para mejor control
CONCURRENT_REQUESTS = 5  # Limita concurrencia
PAUSE_BETWEEN_BATCHES = 2  # Segundos

# Thresholds de calidad
MIN_CONFIDENCE_ACCEPTED = 60.0
MIN_CONFIDENCE_MANUAL_REVIEW = 40.0

# =============================================================================
# VALIDACIÓN POST-PROCESAMIENTO
# =============================================================================

class QualityFilter:
    """Filtros de calidad post-validación"""
    
    # Palabras clave que indican empresa farmacéutica real
    PHARMA_KEYWORDS = [
        # Español
        "farmacéutica", "farmaceutico", "medicamentos", "laboratorio",
        "pharma", "medicinas", "fármacos", "salud", "biotecnología",
        "biotecnologico", "dispositivos médicos", "suplementos",
        # Inglés
        "pharmaceutical", "pharmaceuticals", "medicines", "drugs",
        "healthcare", "biotech", "medical devices", "supplements",
    ]
    
    @classmethod
    def validate_result(cls, row: pd.Series) -> Dict[str, any]:
        """
        Valida resultado final
        
        Returns:
            {
                "status": "accepted" | "manual_review" | "rejected",
                "reasons": List[str]
            }
        """
        reasons = []
        
        # 1. Sin sitio web -> Rejected
        if not row.get("sitio_web") or row["sitio_web"] == "":
            return {
                "status": "rejected",
                "reasons": ["Sin sitio web encontrado"]
            }
        
        confidence = float(row.get("website_confidence", 0))
        
        # 2. Confidence muy baja -> Rejected
        if confidence < MIN_CONFIDENCE_MANUAL_REVIEW:
            return {
                "status": "rejected",
                "reasons": [f"Confianza muy baja ({confidence}%)"]
            }
        
        # 3. Confidence media -> Manual Review
        if confidence < MIN_CONFIDENCE_ACCEPTED:
            return {
                "status": "manual_review",
                "reasons": [f"Confianza media ({confidence}%), requiere revisión"]
            }
        
        # 4. Verifica que el título sea relevante para pharma
        title = str(row.get("homepage_title", "")).lower()
        has_pharma_keyword = any(kw in title for kw in cls.PHARMA_KEYWORDS)
        
        if not has_pharma_keyword and confidence < 80:
            reasons.append("Título no parece farmacéutico")
            return {
                "status": "manual_review",
                "reasons": reasons
            }
        
        # 5. Verifica que el dominio no sea sospechoso
        domain = extract_domain(row.get("sitio_web", ""))
        if is_blocked_domain(domain):
            return {
                "status": "rejected",
                "reasons": ["Dominio en blacklist"]
            }
        
        # 6. Todo OK -> Accepted
        reasons.append(f"Confianza alta ({confidence}%)")
        if has_pharma_keyword:
            reasons.append("Título relevante para pharma")
        
        return {
            "status": "accepted",
            "reasons": reasons
        }

# =============================================================================
# PROCESADOR PRINCIPAL
# =============================================================================

class WebsiteEnricher:
    """Enriquecedor de sitios web con validación"""
    
    def __init__(self):
        self.ddg = DDGS()
        self.quality_filter = QualityFilter()
        self.stats = {
            "total": 0,
            "processed": 0,
            "accepted": 0,
            "manual_review": 0,
            "rejected": 0,
            "errors": 0
        }
    
    async def process_batch(
        self,
        session: aiohttp.ClientSession,
        batch: pd.DataFrame
    ) -> List[Dict]:
        """Procesa un lote de empresas"""
        
        # Crea tareas limitadas por semáforo
        semaphore = asyncio.Semaphore(CONCURRENT_REQUESTS)
        
        async def process_one(row):
            async with semaphore:
                return await self._process_company(session, row)
        
        tasks = [process_one(row) for _, row in batch.iterrows()]
        results = await asyncio.gather(*tasks, return_exceptions=True)
        
        # Filtra excepciones
        valid_results = []
        for result in results:
            if isinstance(result, Exception):
                print(f"⚠️ Error en batch: {result}")
                self.stats["errors"] += 1
            else:
                valid_results.append(result)
        
        return valid_results
    
    async def _process_company(
        self,
        session: aiohttp.ClientSession,
        row: pd.Series
    ) -> Dict:
        """Procesa una empresa individual"""
        
        nombre = row.get("nombre_empresa", "")
        pais = row.get("pais", "")
        
        try:
            # 1. Resuelve sitio oficial
            result = await resolve_official_site(
                session=session,
                ddg=self.ddg,
                company_name=nombre,
                country=pais,
                hinted_url=row.get("sitio_web", "") or row.get("fuente", "")
            )
            
            # 2. Combina con datos originales
            enriched = {
                **row.to_dict(),
                **result,
                "last_checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            
            # 3. Valida calidad
            validation = self.quality_filter.validate_result(
                pd.Series(enriched)
            )
            
            enriched["validation_status"] = validation["status"]
            enriched["validation_reasons"] = "; ".join(validation["reasons"])
            
            # 4. Actualiza stats
            self.stats[validation["status"]] += 1
            self.stats["processed"] += 1
            
            return enriched
        
        except Exception as e:
            print(f"❌ Error procesando {nombre}: {e}")
            self.stats["errors"] += 1
            
            return {
                **row.to_dict(),
                "sitio_web": "",
                "website_status": 0,
                "homepage_title": "",
                "website_confidence": 0.0,
                "website_source": "error",
                "validation_status": "rejected",
                "validation_reasons": f"Error: {str(e)[:100]}",
                "validation_reasons_detail": "",
                "last_checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
    
    def print_progress(self, batch_num: int, total_batches: int):
        """Imprime progreso"""
        print(f"\n{'='*80}")
        print(f"📊 Batch {batch_num}/{total_batches} | Progreso Global")
        print(f"{'='*80}")
        print(f"  Total procesadas:    {self.stats['processed']}/{self.stats['total']}")
        print(f"  ✅ Aceptadas:        {self.stats['accepted']} ({self.stats['accepted']/max(1, self.stats['processed'])*100:.1f}%)")
        print(f"  🔍 Revisión manual:  {self.stats['manual_review']} ({self.stats['manual_review']/max(1, self.stats['processed'])*100:.1f}%)")
        print(f"  ❌ Rechazadas:       {self.stats['rejected']} ({self.stats['rejected']/max(1, self.stats['processed'])*100:.1f}%)")
        if self.stats['errors'] > 0:
            print(f"  ⚠️  Errores:          {self.stats['errors']}")
        print(f"{'='*80}\n")
    
    async def enrich_dataframe(self, df: pd.DataFrame) -> pd.DataFrame:
        """Enriquece DataFrame completo"""
        
        self.stats["total"] = len(df)
        print(f"\n🚀 Iniciando enriquecimiento de {len(df)} empresas...")
        print(f"   Confidence mínima para aceptar: {MIN_CONFIDENCE_ACCEPTED}%")
        print(f"   Confidence mínima para revisión: {MIN_CONFIDENCE_MANUAL_REVIEW}%")
        print(f"   Batch size: {BATCH_SIZE}")
        print(f"   Concurrencia: {CONCURRENT_REQUESTS}")
        
        all_results = []
        total_batches = (len(df) + BATCH_SIZE - 1) // BATCH_SIZE
        
        async with aiohttp.ClientSession() as session:
            for i in range(0, len(df), BATCH_SIZE):
                batch_num = (i // BATCH_SIZE) + 1
                batch = df.iloc[i:i+BATCH_SIZE]
                
                print(f"\n🔄 Procesando batch {batch_num}/{total_batches}...")
                
                batch_results = await self.process_batch(session, batch)
                all_results.extend(batch_results)
                
                self.print_progress(batch_num, total_batches)
                
                # Pausa entre batches (evita rate limits)
                if i + BATCH_SIZE < len(df):
                    print(f"⏸️  Pausa de {PAUSE_BETWEEN_BATCHES}s...")
                    await asyncio.sleep(PAUSE_BETWEEN_BATCHES)
        
        return pd.DataFrame(all_results)

# =============================================================================
# EXPORTACIÓN CON FILTROS
# =============================================================================

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

def export_results(df: pd.DataFrame):
    """Exporta resultados simplificados a Excel con hipervínculos"""
    
    # Solo mantenemos las columnas requeridas
    cols = ["nombre_empresa", "pais", "giro", "sitio_web"]
    for c in cols:
        if c not in df.columns:
            df[c] = ""
    df_out = df[cols].copy()

    # Si quieres aplicar filtro (solo aceptadas o revisadas)
    df_out = df_out[df["validation_status"].isin(["accepted", "manual_review"])]

    # Ruta de salida
    output_excel = OUTPUT_FILE.parent / "empresas_sitios.xlsx"

    # Crea workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sitios BePharma"

    # Encabezados
    headers = ["Nombre de Empresa", "País", "Giro", "Sitio Web"]
    ws.append(headers)

    # Estilo de encabezados
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Inserta datos
    for _, row in df_out.iterrows():
        nombre = row["nombre_empresa"]
        pais = row["pais"]
        giro = row["giro"]
        sitio = row["sitio_web"]

        # Si tiene sitio, lo convierte en hipervínculo clicable
        if sitio and isinstance(sitio, str) and sitio.strip():
            cell_val = sitio
        else:
            cell_val = ""

        ws.append([nombre, pais, giro, cell_val])

        # Crea el hipervínculo en la celda de sitio web
        if sitio and sitio.strip():
            cell = ws.cell(row=ws.max_row, column=4)
            cell.hyperlink = sitio if sitio.startswith("http") else f"https://{sitio}"
            cell.font = Font(color="0000EE", underline="single")

    # Ajuste automático de ancho de columnas
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    wb.save(output_excel)
    print(f"\n✅ Archivo Excel generado con hipervínculos en: {output_excel}")

# =============================================================================
# MAIN
# =============================================================================

async def main():
    """Función principal"""
    
    # 1. Carga datos
    print(f"📂 Cargando datos desde: {INPUT_FILE}")
    df = pd.read_csv(INPUT_FILE, encoding="utf-8-sig")
    print(f"   {len(df)} empresas cargadas")
    
    # 2. Validación básica
    required_cols = ["nombre_empresa", "pais"]
    missing_cols = [c for c in required_cols if c not in df.columns]
    if missing_cols:
        raise ValueError(f"Columnas faltantes: {missing_cols}")
    
    # 3. Limpia datos
    df = df[df["nombre_empresa"].notna() & (df["nombre_empresa"] != "")]
    df = df[df["pais"].notna() & (df["pais"] != "")]
    print(f"   {len(df)} empresas después de limpieza")
    
    # 4. Enriquece
    enricher = WebsiteEnricher()
    enriched_df = await enricher.enrich_dataframe(df)
    
    # 5. Exporta
    export_results(enriched_df)
    
    print("\n✅ Proceso completado exitosamente")

if __name__ == "__main__":
    asyncio.run(main())